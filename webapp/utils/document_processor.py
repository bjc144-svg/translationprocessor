"""
Document Processor
Handles Word document processing, header/footer addition, certificate generation, and PDF conversion
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from pathlib import Path
import tempfile
import shutil


class DocumentProcessor:
    def __init__(self):
        """Initialize document processor"""

        # Get templates directory
        self.templates_dir = Path(__file__).parent.parent / 'templates'
        self.templates_dir.mkdir(parents=True, exist_ok=True)

        # Get assets directory
        self.assets_dir = Path(__file__).parent.parent / 'assets'
        self.assets_dir.mkdir(parents=True, exist_ok=True)

    def convert_doc_to_docx(self, doc_path):
        """Convert .doc file to .docx using Word COM"""
        try:
            import pythoncom
            import win32com.client

            # Initialize COM
            pythoncom.CoInitialize()

            try:
                # Create output path
                docx_path = str(doc_path).replace('.doc', '.docx')

                # Open Word
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False

                # Open the .doc file
                doc = word.Documents.Open(str(Path(doc_path).absolute()))

                # Save as .docx (FileFormat 16 = .docx)
                doc.SaveAs2(str(Path(docx_path).absolute()), FileFormat=16)
                doc.Close()
                word.Quit()

                print(f"Converted .doc to .docx: {docx_path}")
                return docx_path

            finally:
                try:
                    pythoncom.CoUninitialize()
                except:
                    pass

        except Exception as e:
            print(f"Error converting .doc to .docx: {e}")
            import traceback
            traceback.print_exc()
            return None

    def process_translation(self, input_file, output_file, metadata):
        """
        Process translation document

        Args:
            input_file: Path to input Word document
            output_file: Path to output PDF file
            metadata: Dictionary containing:
                - case_number
                - language_pair
                - source_language
                - target_language
                - translator_name
                - translator_signature (path to signature image)
                - date

        Returns:
            True if successful, False otherwise
        """

        try:
            print("Starting document processing...")
            print(f"Input: {input_file}")
            print(f"Output: {output_file}")
            print(f"Metadata: {metadata}")

            # Convert .doc to .docx if needed
            input_path = Path(input_file)
            if input_path.suffix.lower() == '.doc':
                print("Detected .doc file, converting to .docx...")
                converted_path = self.convert_doc_to_docx(input_file)
                if converted_path:
                    input_file = converted_path
                else:
                    print("Error: Failed to convert .doc to .docx")
                    return False

            # Create temporary directory for processing
            with tempfile.TemporaryDirectory() as temp_dir:
                temp_path = Path(temp_dir)

                # Step 1: Add header and footer to translation document
                print("Step 1: Adding header and footer to translation...")
                translation_docx = temp_path / "translation_with_header.docx"
                self.add_header_footer(input_file, translation_docx, metadata)

                # Determine which certificates to include based on division
                division = metadata.get('division', 'PARK')
                include_translator_cert = (division != 'EEI')

                documents_to_combine = [translation_docx]

                # Step 2: Create translator certificate (skip if division is EEI)
                if include_translator_cert:
                    print("Step 2: Creating translator certificate...")
                    translator_cert_docx = temp_path / "translator_certificate.docx"
                    self.create_translator_certificate(translator_cert_docx, metadata)
                    documents_to_combine.append(translator_cert_docx)
                else:
                    print("Step 2: Skipping translator certificate (division is EEI)...")

                # Step 3: Create Park certificate
                print("Step 3: Creating Park certificate...")
                park_cert_docx = temp_path / "park_certificate.docx"
                self.create_park_certificate(park_cert_docx, metadata)
                documents_to_combine.append(park_cert_docx)

                # Step 4: Combine documents
                print(f"Step 4: Combining documents ({len(documents_to_combine)} total)...")
                combined_docx = temp_path / "combined.docx"
                self.combine_documents(
                    documents_to_combine,
                    combined_docx
                )

                # Step 5: Convert to PDF
                print("Step 5: Converting to PDF...")
                self.convert_to_pdf(combined_docx, output_file)

                print("Processing complete!")
                return True

        except Exception as e:
            print(f"Error processing document: {e}")
            import traceback
            traceback.print_exc()
            return False

    def _find_libreoffice(self):
        """
        Find LibreOffice executable path (cross-platform).

        Returns:
            Path to LibreOffice executable, or None if not found
        """
        import platform
        import subprocess
        from pathlib import Path

        system = platform.system()

        if system == 'Windows':
            # Common LibreOffice installation paths on Windows
            possible_paths = [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                Path(os.environ.get('PROGRAMFILES', 'C:\\Program Files')) / 'LibreOffice' / 'program' / 'soffice.exe',
                Path(os.environ.get('PROGRAMFILES(X86)', 'C:\\Program Files (x86)')) / 'LibreOffice' / 'program' / 'soffice.exe',
            ]

            for path in possible_paths:
                if isinstance(path, str):
                    path = Path(path)
                if path.exists():
                    print(f"  Found LibreOffice at: {path}")
                    return str(path)

            return None

        elif system == 'Linux' or system == 'Darwin':  # Darwin is macOS
            # On Linux/Mac, try to find in PATH
            try:
                result = subprocess.run(['which', 'libreoffice'],
                                      capture_output=True, text=True, timeout=5)
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.strip()

                # Alternative name on some systems
                result = subprocess.run(['which', 'soffice'],
                                      capture_output=True, text=True, timeout=5)
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.strip()
            except:
                pass

            # Try common paths
            common_paths = ['/usr/bin/libreoffice', '/usr/bin/soffice',
                          '/usr/local/bin/libreoffice', '/usr/local/bin/soffice']
            for path in common_paths:
                if Path(path).exists():
                    return path

            return None

        return None

    def _find_poppler(self):
        """
        Find poppler binaries path (cross-platform).

        Returns:
            Path to poppler bin directory, or None if not found
        """
        import platform
        from pathlib import Path

        system = platform.system()

        if system == 'Windows':
            # Common poppler installation paths on Windows
            possible_paths = [
                r"C:\Program Files\poppler\Library\bin",
                r"C:\Program Files (x86)\poppler\Library\bin",
                r"C:\poppler\Library\bin",
                r"C:\Program Files\poppler-24.02.0\Library\bin",  # Version-specific
                r"C:\Program Files\poppler-23.11.0\Library\bin",
                Path(os.environ.get('PROGRAMFILES', 'C:\\Program Files')) / 'poppler' / 'Library' / 'bin',
                Path(os.environ.get('LOCALAPPDATA', 'C:\\Users\\' + os.environ.get('USERNAME', 'User') + '\\AppData\\Local')) / 'poppler' / 'Library' / 'bin',
            ]

            # Also check if poppler is in PATH
            import subprocess
            try:
                result = subprocess.run(['where', 'pdftoppm.exe'],
                                      capture_output=True, text=True, timeout=5)
                if result.returncode == 0 and result.stdout.strip():
                    # Get the directory containing pdftoppm.exe
                    pdftoppm_path = result.stdout.strip().split('\n')[0]
                    return str(Path(pdftoppm_path).parent)
            except:
                pass

            for path in possible_paths:
                if isinstance(path, str):
                    path = Path(path)
                if path.exists() and (path / 'pdftoppm.exe').exists():
                    print(f"  Found poppler at: {path}")
                    return str(path)

            return None

        elif system == 'Linux' or system == 'Darwin':
            # On Linux/Mac, poppler is usually in PATH
            import subprocess
            try:
                result = subprocess.run(['which', 'pdftoppm'],
                                      capture_output=True, text=True, timeout=5)
                if result.returncode == 0 and result.stdout.strip():
                    # Return the directory, not the full path to pdftoppm
                    return str(Path(result.stdout.strip()).parent)
            except:
                pass

            # Try common paths
            common_paths = ['/usr/bin', '/usr/local/bin']
            for path in common_paths:
                if Path(path).exists() and (Path(path) / 'pdftoppm').exists():
                    return path

            return None

        return None

    def _convert_docx_to_images(self, docx_path, dpi=150):
        """
        Convert .docx pages to images using LibreOffice and pdf2image.

        Args:
            docx_path: Path to .docx file
            dpi: Resolution for images (default 150 for good quality/size balance)

        Returns:
            List of PIL Image objects, one per page
        """
        import subprocess
        from pdf2image import convert_from_path
        from pathlib import Path

        print(f"Converting {docx_path} to images...")

        # Find LibreOffice executable
        libreoffice_path = self._find_libreoffice()
        if not libreoffice_path:
            raise Exception(
                "LibreOffice not found. Please install LibreOffice from https://www.libreoffice.org/download/"
            )

        # Find poppler binaries
        poppler_path = self._find_poppler()
        if not poppler_path:
            import platform
            if platform.system() == 'Windows':
                raise Exception(
                    "Poppler not found. Please install poppler for Windows:\n"
                    "1. Download from: https://github.com/oschwartz10612/poppler-windows/releases/\n"
                    "2. Extract to C:\\Program Files\\poppler\\ (or any location)\n"
                    "3. The bin folder should contain pdftoppm.exe and pdfinfo.exe"
                )
            else:
                raise Exception(
                    "Poppler not found. Please install poppler:\n"
                    "- Ubuntu/Debian: sudo apt-get install poppler-utils\n"
                    "- macOS: brew install poppler"
                )

        # Create temporary directory for conversion
        temp_dir = Path(tempfile.mkdtemp(prefix='docx2img_'))

        try:
            # Step 1: Convert .docx to PDF using LibreOffice
            pdf_path = temp_dir / "temp.pdf"

            print(f"  Step 1: Converting .docx to PDF using LibreOffice...")
            result = subprocess.run([
                libreoffice_path,
                '--headless',  # Run without GUI
                '--convert-to', 'pdf',
                '--outdir', str(temp_dir),
                str(docx_path)
            ], capture_output=True, text=True, timeout=60)

            if result.returncode != 0:
                raise Exception(f"LibreOffice conversion failed: {result.stderr}")

            # LibreOffice names the output file based on input filename
            docx_filename = Path(docx_path).stem
            pdf_path = temp_dir / f"{docx_filename}.pdf"

            if not pdf_path.exists():
                raise Exception(f"PDF not created at expected path: {pdf_path}")

            print(f"  Step 2: Converting PDF pages to images (DPI={dpi})...")
            # Step 2: Convert PDF pages to images
            images = convert_from_path(
                str(pdf_path),
                dpi=dpi,
                fmt='png',  # PNG for better quality
                thread_count=2,  # Use multiple threads for faster conversion
                poppler_path=poppler_path  # Provide poppler path for Windows
            )

            print(f"  ✓ Converted {len(images)} pages to images")
            return images

        except subprocess.TimeoutExpired:
            raise Exception("LibreOffice conversion timed out after 60 seconds")
        except Exception as e:
            print(f"  ✗ Error converting to images: {e}")
            raise
        finally:
            # Clean up temporary directory
            try:
                shutil.rmtree(temp_dir)
            except Exception as e:
                print(f"Warning: Could not clean up temp directory {temp_dir}: {e}")

    def add_header_footer(self, input_file, output_file, metadata):
        """Add header and footer to translation document using page-to-image conversion"""

        print("Converting source document pages to images for perfect visual preservation...")

        # Convert source document pages to images
        page_images = self._convert_docx_to_images(input_file, dpi=150)
        total_pages = len(page_images)

        print(f"Converted {total_pages} pages to images. Creating new document with headers/footers...")

        # Create a new document
        doc = Document()

        # Remove the initial empty paragraph that Document() creates
        if doc.paragraphs:
            for para in list(doc.paragraphs):
                p_element = para._element
                p_element.getparent().remove(p_element)

        # Get page setup from original document for each page/section
        original_doc = Document(input_file)
        print(f"Original document has {len(original_doc.sections)} section(s)")

        # Build a list of page properties from source document
        # Each source section may contain multiple pages, but we need per-page info
        # We'll map each output page to the appropriate source section
        source_sections = list(original_doc.sections)

        # Add each page image to the document
        from io import BytesIO

        for page_num, page_image in enumerate(page_images):
            print(f"  Adding page {page_num + 1}/{total_pages} as image...")

            # Determine which source section this page belongs to
            # For simplicity, we'll map linearly if we have fewer sections than pages
            source_section_idx = min(page_num, len(source_sections) - 1) if source_sections else 0
            source_section = source_sections[source_section_idx] if source_sections else None

            # Save image to bytes
            img_bytes = BytesIO()
            page_image.save(img_bytes, format='PNG')
            img_bytes.seek(0)

            # Add paragraph with full-page image
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run()

            # Calculate image dimensions to fit within available space
            # We need to account for margins AND header/footer space
            if source_section:
                page_width = source_section.page_width
                page_height = source_section.page_height
                left_margin = source_section.left_margin
                right_margin = source_section.right_margin
                top_margin = source_section.top_margin
                bottom_margin = source_section.bottom_margin

                # Calculate available content area
                available_width = page_width - left_margin - right_margin

                # Header/footer distances (will be set to 0.15" later)
                header_distance = Inches(0.15)
                footer_distance = Inches(0.15)

                # Estimate header/footer heights (conservative estimates)
                # Header: logo (0.4") + table spacing + line = ~0.6"
                # Footer: line + title + table = ~0.5"
                header_height = Inches(0.6)
                footer_height = Inches(0.5)

                # Available height for content (accounting for header/footer)
                available_height = (page_height - top_margin - bottom_margin -
                                  header_distance - footer_distance -
                                  header_height - footer_height)

                # Get the actual image dimensions
                img_width_px, img_height_px = page_image.size
                img_aspect_ratio = img_height_px / img_width_px

                # Calculate image size to fit within available space
                # Try fitting by width first
                img_width = available_width
                img_height = img_width * img_aspect_ratio

                # If height exceeds available space, scale down
                if img_height > available_height:
                    img_height = available_height
                    img_width = img_height / img_aspect_ratio

                print(f"    Page {page_num + 1}: {page_width/Inches(1):.2f}\" x {page_height/Inches(1):.2f}\", "
                      f"image: {img_width/Inches(1):.2f}\" x {img_height/Inches(1):.2f}\"")
            else:
                # Default to letter size
                img_width = Inches(6.5)

            # Add the image with calculated dimensions
            run.add_picture(img_bytes, width=img_width)

            # Add section break after this page (except for the last page)
            if page_num < total_pages - 1:
                # Add section break to start a new page
                para_elem = para._element
                pPr = para_elem.get_or_add_pPr()
                sectPr = OxmlElement('w:sectPr')
                pPr.append(sectPr)

                # Set section break type to nextPage
                sectType = OxmlElement('w:type')
                sectType.set(qn('w:val'), 'nextPage')
                sectPr.append(sectType)

        print(f"Added {total_pages} page images to document")
        print(f"Document now has {len(doc.sections)} section(s)")

        # Copy page setup from source document on a per-section basis
        print("Copying page setup from source document (per-page)...")
        for section_idx, section in enumerate(doc.sections):
            # Map each output section to corresponding source section
            source_section_idx = min(section_idx, len(source_sections) - 1) if source_sections else 0
            source_section = source_sections[source_section_idx] if source_sections else None

            if source_section:
                try:
                    section.top_margin = source_section.top_margin
                    section.bottom_margin = source_section.bottom_margin
                    section.left_margin = source_section.left_margin
                    section.right_margin = source_section.right_margin
                    section.page_height = source_section.page_height
                    section.page_width = source_section.page_width
                    section.orientation = source_section.orientation
                    print(f"  Section {section_idx}: {section.page_width/Inches(1):.2f}\" x {section.page_height/Inches(1):.2f}\", orientation: {section.orientation}")
                except Exception as e:
                    print(f"Warning: Could not copy page setup for section {section_idx}: {e}")

        # Apply header/footer to ALL sections
        print("Adding headers and footers to all sections...")
        for section_idx, section in enumerate(doc.sections):
            # Minimize header/footer distances to maximize content area
            section.header_distance = Inches(0.15)  # Minimal distance from top
            section.footer_distance = Inches(0.15)  # Minimal distance from bottom

            # Set up header
            header = section.header
            header.is_linked_to_previous = False

            # Clear existing header
            for paragraph in header.paragraphs:
                paragraph.clear()

            # Create header table (2 columns) - span full width
            header_table = header.add_table(1, 2, Inches(6.5))
            header_table.autofit = False
            header_table.allow_autofit = False

            # Set table to full width with no spacing
            tbl = header_table._element
            tblPr = tbl.tblPr
            if tblPr is None:
                tblPr = OxmlElement('w:tblPr')
                tbl.insert(0, tblPr)

            # Set table width to 100% (5000 = 100% in Word's measurement)
            tblW = OxmlElement('w:tblW')
            tblW.set(qn('w:w'), '5000')
            tblW.set(qn('w:type'), 'pct')
            tblPr.append(tblW)

            # Remove table cell spacing
            tblCellSpacing = OxmlElement('w:tblCellSpacing')
            tblCellSpacing.set(qn('w:w'), '0')
            tblCellSpacing.set(qn('w:type'), 'dxa')
            tblPr.append(tblCellSpacing)

            # Set column widths (proportional)
            header_table.rows[0].cells[0].width = Inches(2.75)
            header_table.rows[0].cells[1].width = Inches(3.75)

            # Left cell - Logo
            left_cell = header_table.rows[0].cells[0]
            left_para = left_cell.paragraphs[0]
            left_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # Minimize paragraph spacing
            left_para.paragraph_format.space_before = Pt(0)
            left_para.paragraph_format.space_after = Pt(0)
            left_para.paragraph_format.line_spacing = 1.0

            # Try to add logo image (use EEI logo for EEI division, Park logo otherwise)
            division = metadata.get('division', 'PARK')
            if division == 'EEI':
                logo_path = self.assets_dir / 'eei_logo.png'
                logo_name = "EEI"
            else:
                logo_path = self.assets_dir / 'park_logo.png'
                logo_name = "Park"

            logo_added = False

            if logo_path.exists():
                try:
                    run = left_para.add_run()
                    run.add_picture(str(logo_path), height=Inches(0.4))  # Reduced from 0.5
                    logo_added = True
                    print(f"Successfully added {logo_name} logo from: {logo_path}")
                except Exception as e:
                    print(f"Error adding logo image: {e}")
                    import traceback
                    traceback.print_exc()

            # If logo couldn't be loaded, use text fallback
            if not logo_added:
                fallback_text = "PARK EVALUATION SERVICES" if division != 'EEI' else "EEI"
                left_para.text = fallback_text
                left_para.runs[0].font.bold = True
                left_para.runs[0].font.size = Pt(9)  # Reduced from 10
                print(f"Warning: {logo_name} logo file not found at: {logo_path}")

            # Right cell - Contact info
            right_cell = header_table.rows[0].cells[1]
            right_cell.vertical_alignment = 1  # Center vertically

            # Remove cell margins for proper right alignment
            tc = right_cell._element
            tcPr = tc.get_or_add_tcPr()
            tcMar = OxmlElement('w:tcMar')
            for margin_name in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement(f'w:{margin_name}')
                node.set(qn('w:w'), '0')
                node.set(qn('w:type'), 'dxa')
                tcMar.append(node)
            tcPr.append(tcMar)

            # Add contact info (right-aligned)
            contact_lines = [
                "212-581-8877",
                "eval@parkeval.com",
                "www.parkeval.com"
            ]

            for i, line in enumerate(contact_lines):
                if i > 0:
                    right_cell.add_paragraph()
                para = right_cell.paragraphs[i] if i == 0 else right_cell.paragraphs[-1]
                para.text = line
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                # Remove paragraph spacing for tight alignment
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                run = para.runs[0]
                run.font.size = Pt(9)  # Reduced from 10

            # Add horizontal line after header using border
            line_para = header.add_paragraph()
            pPr = line_para._element.get_or_add_pPr()
            # Remove paragraph spacing
            spacing = OxmlElement('w:spacing')
            spacing.set(qn('w:before'), '0')
            spacing.set(qn('w:after'), '0')
            pPr.append(spacing)
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '0')
            bottom.set(qn('w:color'), '000000')
            pBdr.append(bottom)
            pPr.append(pBdr)

            # Set up footer
            footer = section.footer
            footer.is_linked_to_previous = False

            # Clear existing footer
            for paragraph in footer.paragraphs:
                paragraph.clear()

            # Add horizontal line before footer
            line_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            pPr = line_para._element.get_or_add_pPr()
            # Remove paragraph spacing
            spacing = OxmlElement('w:spacing')
            spacing.set(qn('w:before'), '0')
            spacing.set(qn('w:after'), '0')
            pPr.append(spacing)
            pBdr = OxmlElement('w:pBdr')
            top = OxmlElement('w:top')
            top.set(qn('w:val'), 'single')
            top.set(qn('w:sz'), '6')
            top.set(qn('w:space'), '0')
            top.set(qn('w:color'), '000000')
            pBdr.append(top)
            pPr.append(pBdr)

            # Add "CERTIFIED TRANSLATION" (centered, all caps, not bold)
            cert_para = footer.add_paragraph("CERTIFIED TRANSLATION")
            cert_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Make spacing extremely tight to minimize footer height
            cert_para.paragraph_format.space_before = Pt(0)
            cert_para.paragraph_format.space_after = Pt(0)
            cert_para.paragraph_format.line_spacing = 1.0
            cert_para.runs[0].font.size = Pt(10)  # Reduced from 12

            # Add footer info table (3 columns) - span full width like header
            footer_table = footer.add_table(1, 3, Inches(6.5))
            footer_table.autofit = False
            footer_table.allow_autofit = False

            # Set table to full width with no spacing
            tbl = footer_table._element
            tblPr = tbl.tblPr
            if tblPr is None:
                tblPr = OxmlElement('w:tblPr')
                tbl.insert(0, tblPr)

            # Set table width to 100% (5000 = 100% in Word's measurement)
            tblW = OxmlElement('w:tblW')
            tblW.set(qn('w:w'), '5000')
            tblW.set(qn('w:type'), 'pct')
            tblPr.append(tblW)

            # Remove table cell spacing
            tblCellSpacing = OxmlElement('w:tblCellSpacing')
            tblCellSpacing.set(qn('w:w'), '0')
            tblCellSpacing.set(qn('w:type'), 'dxa')
            tblPr.append(tblCellSpacing)

            # Left - Case number (use EEI or Park based on division)
            left_cell = footer_table.rows[0].cells[0]
            left_para = left_cell.paragraphs[0]
            division = metadata.get('division', 'PARK')
            case_prefix = "EEI" if division == 'EEI' else "Park"
            left_para.text = f"{case_prefix} Case #{metadata['case_number']}"
            left_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            left_para.paragraph_format.space_before = Pt(0)
            left_para.paragraph_format.space_after = Pt(0)
            left_para.runs[0].font.size = Pt(8)  # Reduced from 9

            # Center - Page number with field codes
            center_cell = footer_table.rows[0].cells[1]
            center_para = center_cell.paragraphs[0]
            center_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            center_para.paragraph_format.space_before = Pt(0)
            center_para.paragraph_format.space_after = Pt(0)

            # Add "Page " text
            run = center_para.add_run("Page ")
            run.font.size = Pt(8)  # Reduced from 9

            # Add PAGE field (current page number)
            self._add_page_number(center_para)

            # Add " of " text
            run = center_para.add_run(" of ")
            run.font.size = Pt(8)  # Reduced from 9

            # Add NUMPAGES field (total pages)
            # Calculate how many certificate pages to exclude
            division = metadata.get('division', 'PARK')
            num_cert_pages = 1 if division == 'EEI' else 2  # EEI: Park only, Others: Translator + Park
            self._add_num_pages(center_para, num_cert_pages)

            # Right - Language pair
            right_cell = footer_table.rows[0].cells[2]
            right_para = right_cell.paragraphs[0]
            right_para.text = metadata['language_pair']
            right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            right_para.paragraph_format.space_before = Pt(0)
            right_para.paragraph_format.space_after = Pt(0)
            right_para.runs[0].font.size = Pt(8)  # Reduced from 9

        # Save document
        doc.save(output_file)
        print(f"Saved document with header/footer to: {output_file}")

        # Also save a copy to the output folder for inspection
        import shutil
        from pathlib import Path
        # Get the webapp output directory
        webapp_dir = Path(__file__).parent.parent
        output_folder = webapp_dir / "output"
        output_folder.mkdir(parents=True, exist_ok=True)
        debug_copy = output_folder / "DEBUG_translation_with_header.docx"
        try:
            shutil.copy2(output_file, debug_copy)
            print(f"DEBUG: Saved copy for inspection at: {debug_copy}")
            print(f"DEBUG: Please open this file in Word to check if blank page 2 exists there")
        except Exception as e:
            print(f"DEBUG: Failed to save debug copy: {e}")

    def _add_page_number(self, paragraph):
        """Add PAGE field to paragraph for current page number"""
        run = paragraph.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')

        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = 'PAGE'

        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run.font.size = Pt(8)  # Reduced from 9

    def _add_num_pages(self, paragraph, pages_to_exclude=2):
        """Add formula field to calculate NUMPAGES - N (excludes certificate pages)

        Args:
            paragraph: Paragraph to add the field to
            pages_to_exclude: Number of certificate pages to exclude (1 for EEI, 2 for others)
        """
        run = paragraph.add_run()

        # Begin formula field
        fldChar_formula_begin = OxmlElement('w:fldChar')
        fldChar_formula_begin.set(qn('w:fldCharType'), 'begin')
        run._r.append(fldChar_formula_begin)

        # Formula instruction start: "= "
        instrText_formula_start = OxmlElement('w:instrText')
        instrText_formula_start.set(qn('xml:space'), 'preserve')
        instrText_formula_start.text = '= '
        run._r.append(instrText_formula_start)

        # Begin nested NUMPAGES field
        fldChar_numpages_begin = OxmlElement('w:fldChar')
        fldChar_numpages_begin.set(qn('w:fldCharType'), 'begin')
        run._r.append(fldChar_numpages_begin)

        # NUMPAGES instruction
        instrText_numpages = OxmlElement('w:instrText')
        instrText_numpages.set(qn('xml:space'), 'preserve')
        instrText_numpages.text = 'NUMPAGES'
        run._r.append(instrText_numpages)

        # End nested NUMPAGES field
        fldChar_numpages_end = OxmlElement('w:fldChar')
        fldChar_numpages_end.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar_numpages_end)

        # Formula instruction end: " - N" (where N is pages_to_exclude)
        instrText_formula_end = OxmlElement('w:instrText')
        instrText_formula_end.set(qn('xml:space'), 'preserve')
        instrText_formula_end.text = f' - {pages_to_exclude}'
        run._r.append(instrText_formula_end)

        # End formula field
        fldChar_formula_end = OxmlElement('w:fldChar')
        fldChar_formula_end.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar_formula_end)

        run.font.size = Pt(8)  # Reduced from 9

    def _copy_document_body(self, source_doc, target_doc):
        """
        Copy all body elements from source to target document using deep copy.
        This preserves ALL element types including paragraphs, tables, textboxes,
        SDTs, and other complex structures that don't have python-docx API wrappers.
        """
        from copy import deepcopy
        from io import BytesIO

        # Track relationship ID mappings for images and embedded objects
        rel_mapping = {}

        # Statistics for logging
        stats = {'total': 0, 'deepcopy_success': 0, 'fallback': 0, 'failed': 0}

        for element in source_doc.element.body:
            stats['total'] += 1
            element_tag = element.tag.split('}')[-1]  # Extract tag name without namespace

            # DON'T skip sectPr anymore - we need to preserve section structure
            # Previously we skipped it, but now we need it to define section boundaries
            # if element_tag == 'sectPr':
            #     print(f"⊘ Skipping sectPr (section properties) - using target document's section settings")
            #     stats['total'] -= 1  # Don't count as processed
            #     continue

            try:
                # Attempt deep copy (preserves ALL XML structure)
                new_element = deepcopy(element)

                # Reconcile relationship IDs (images, shapes, embedded objects)
                self._reconcile_relationships(new_element, source_doc, target_doc, rel_mapping)

                # DON'T remove section properties - we need to preserve them for different page setups
                # Different pages may have different sizes/orientations (e.g., A4 portrait, then A4 landscape)
                # self._remove_section_properties(new_element)  # Commented out to preserve sections

                # Append to target document body
                # We've removed the target's initial sectPr, so we can append normally
                # Section breaks from source (sectPr elements) will create proper section structure
                target_doc.element.body.append(new_element)

                stats['deepcopy_success'] += 1
                print(f"✓ Deep copied element: {element_tag}")

            except Exception as e:
                print(f"⚠ Deep copy failed for {element_tag}: {e}")
                import traceback
                traceback.print_exc()

                # Try fallback handlers for known element types
                if element_tag == 'p':
                    if self._fallback_copy_paragraph(element, source_doc, target_doc):
                        stats['fallback'] += 1
                    else:
                        stats['failed'] += 1
                elif element_tag == 'tbl':
                    if self._fallback_copy_table(element, source_doc, target_doc):
                        stats['fallback'] += 1
                    else:
                        stats['failed'] += 1
                else:
                    # Unknown element type with no fallback
                    print(f"✗ ERROR: No fallback handler for {element_tag}")
                    stats['failed'] += 1

        # Log final statistics
        print(f"Body copy complete: {stats}")
        if stats['failed'] > 0:
            print(f"WARNING: {stats['failed']} elements could not be copied")

    def _reconcile_relationships(self, element, source_doc, target_doc, rel_map):
        """
        Walk copied element tree and remap all relationship IDs.
        Handles images, embedded objects, and other linked resources.
        """
        rel_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

        # Attributes that contain relationship IDs
        rel_attrs = [
            f'{rel_ns}embed',  # Embedded images/objects
            f'{rel_ns}link',   # Linked images/objects
            f'{rel_ns}id',     # General relationships
        ]

        # Walk all descendants of the element
        # We need to collect elements to remove so we don't modify while iterating
        elements_to_remove = []

        for child in element.iter():
            for attr in rel_attrs:
                old_rid = child.get(attr)
                if old_rid and old_rid in source_doc.part.related_parts:
                    # Get or create new relationship
                    new_rid = self._copy_relationship(old_rid, source_doc, target_doc, rel_map)

                    if new_rid is None:
                        # For headerReference and footerReference elements, remove the entire element
                        # Just deleting the r:id attribute leaves an invalid element
                        tag_name = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                        if tag_name in ('headerReference', 'footerReference'):
                            elements_to_remove.append((child.getparent(), child))
                        else:
                            # For other elements, just delete the attribute
                            if attr in child.attrib:
                                del child.attrib[attr]
                    else:
                        # Update element to use new relationship ID
                        child.set(attr, new_rid)

        # Remove collected elements
        for parent, child in elements_to_remove:
            if parent is not None:
                parent.remove(child)

    def _remove_section_properties(self, element):
        """
        Remove all sectPr (section properties) elements from within the element tree.
        Section breaks from the source document would create unwanted sections in target.
        We want all content in a single section with our headers/footers.
        """
        # Find all sectPr elements within the element tree
        # These can appear in paragraph properties (pPr)
        w_ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        sectPr_tag = f'{w_ns}sectPr'

        # Collect all sectPr elements to remove
        sectPr_elements = []
        for child in element.iter():
            if child.tag == sectPr_tag:
                sectPr_elements.append((child.getparent(), child))

        # Remove them
        for parent, sectPr in sectPr_elements:
            if parent is not None:
                parent.remove(sectPr)
                print(f"  ⊘ Removed sectPr from within element (prevents section breaks)")

    def _copy_relationship(self, old_rid, source_doc, target_doc, rel_map):
        """
        Copy a relationship and its associated part to target document.
        Returns new relationship ID, or None if the relationship should be removed.
        """
        # Check cache first
        if old_rid in rel_map:
            return rel_map[old_rid]

        # Get source relationship and part
        if old_rid not in source_doc.part.related_parts:
            print(f"Warning: Relationship {old_rid} not found in source document")
            return old_rid  # Return unchanged

        source_part = source_doc.part.related_parts[old_rid]
        part_data = source_part.blob
        content_type = source_part.content_type

        # Handle images (most common case)
        if 'image' in content_type:
            try:
                from io import BytesIO

                # Use the high-level API to add the image (which handles all the complexity)
                # Create a temporary paragraph, add picture to it, extract the rId, then delete it
                temp_para = target_doc.add_paragraph()
                temp_run = temp_para.add_run()

                # Add the picture using the proven working API
                image_stream = BytesIO(part_data)
                temp_run.add_picture(image_stream)

                # Extract the relationship ID that was just created
                # The picture element contains the relationship reference
                pic_elements = temp_run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/picture}pic')
                if pic_elements:
                    blip = pic_elements[0].find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                    if blip is not None:
                        new_rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')

                        # Delete the temporary paragraph
                        target_doc._element.body.remove(temp_para._element)

                        # Cache mapping
                        rel_map[old_rid] = new_rid
                        print(f"  Mapped image relationship: {old_rid} → {new_rid}")
                        return new_rid

                # If we couldn't extract the rId, delete temp paragraph and return None
                target_doc._element.body.remove(temp_para._element)
                print(f"Warning: Could not extract relationship ID for image {old_rid}")
                return None

            except Exception as e:
                print(f"Error copying image relationship {old_rid}: {e}")
                import traceback
                traceback.print_exc()
                return None  # Signal to remove this relationship reference

        # Skip header/footer relationships - we add our own headers/footers
        # These relationship IDs from the source document don't exist in the target
        if 'header' in content_type or 'footer' in content_type:
            print(f"⊘ Skipping {content_type} relationship (we add our own headers/footers)")
            return None  # Signal to remove this relationship reference

        # For other types (embedded objects, charts, etc.), skip them for now
        # Can be expanded later if needed
        print(f"⊘ Skipping relationship type '{content_type}' (not yet supported)")
        return None  # Signal to remove this relationship reference

    def _fallback_copy_paragraph(self, element, source_doc, target_doc):
        """
        Fallback method to copy paragraph when deep copy fails.
        Uses the original paragraph-by-paragraph logic.
        """
        from io import BytesIO

        try:
            # Create corresponding paragraph in target document
            new_para = target_doc.add_paragraph()

            # Find the original paragraph object
            for para in source_doc.paragraphs:
                if para._element == element:
                    # Copy paragraph properties
                    new_para.alignment = para.alignment
                    new_para.style = para.style

                    # Copy paragraph formatting
                    if para.paragraph_format.space_before:
                        new_para.paragraph_format.space_before = para.paragraph_format.space_before
                    if para.paragraph_format.space_after:
                        new_para.paragraph_format.space_after = para.paragraph_format.space_after
                    if para.paragraph_format.line_spacing:
                        new_para.paragraph_format.line_spacing = para.paragraph_format.line_spacing

                    # Copy runs (text with formatting and images)
                    for run in para.runs:
                        # Check if run contains an image
                        has_image = False
                        drawings = run._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing',
                                                       run._element.nsmap if hasattr(run._element, 'nsmap') else None)

                        if not drawings:
                            for child in run._element:
                                if 'drawing' in child.tag or 'pict' in child.tag:
                                    drawings = [child]
                                    break

                        if drawings:
                            has_image = True
                            # Copy images
                            for drawing in drawings:
                                try:
                                    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                                          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
                                    blip = drawing.find('.//a:blip', ns)

                                    if blip is None:
                                        for elem in drawing.iter():
                                            if 'blip' in elem.tag.lower():
                                                blip = elem
                                                break

                                    if blip is not None:
                                        rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                        if not rId:
                                            rId = blip.get('embed')

                                        if rId and rId in source_doc.part.related_parts:
                                            image_part = source_doc.part.related_parts[rId]
                                            image_bytes = image_part.blob

                                            # Get image dimensions
                                            extent = drawing.find('.//wp:extent',
                                                                {'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'})
                                            if extent is None:
                                                for elem in drawing.iter():
                                                    if 'extent' in elem.tag.lower():
                                                        extent = elem
                                                        break

                                            new_run = new_para.add_run()
                                            if extent is not None and extent.get('cx'):
                                                try:
                                                    cx = int(extent.get('cx'))
                                                    width_inches = cx / 914400.0
                                                    new_run.add_picture(BytesIO(image_bytes), width=Inches(width_inches))
                                                except Exception as ex:
                                                    print(f"Error sizing image: {ex}")
                                                    new_run.add_picture(BytesIO(image_bytes))
                                            else:
                                                new_run.add_picture(BytesIO(image_bytes))
                                except Exception as e:
                                    print(f"Warning: Could not copy image in fallback: {e}")

                        # Copy text if no image
                        if not has_image:
                            new_run = new_para.add_run(run.text)
                            new_run.bold = run.bold
                            new_run.italic = run.italic
                            new_run.underline = run.underline
                            if run.font.size:
                                new_run.font.size = run.font.size
                            if run.font.name:
                                new_run.font.name = run.font.name
                    return True  # Success

            return False  # Paragraph not found

        except Exception as e:
            print(f"Fallback paragraph copy failed: {e}")
            return False

    def _fallback_copy_table(self, element, source_doc, target_doc):
        """
        Fallback method to copy table when deep copy fails.
        Uses the original table-by-table logic.
        """
        try:
            # Find the original table object
            for table in source_doc.tables:
                if table._element == element:
                    # Create table with same dimensions
                    new_table = target_doc.add_table(rows=len(table.rows), cols=len(table.columns))

                    # Copy table style if it exists
                    if table.style:
                        new_table.style = table.style

                    # Copy cell content and formatting
                    for i, row in enumerate(table.rows):
                        for j, cell in enumerate(row.cells):
                            new_cell = new_table.rows[i].cells[j]

                            # Copy each paragraph in the cell
                            new_cell.text = ""  # Clear default paragraph
                            for k, para in enumerate(cell.paragraphs):
                                if k == 0:
                                    cell_para = new_cell.paragraphs[0]
                                else:
                                    cell_para = new_cell.add_paragraph()

                                cell_para.alignment = para.alignment

                                # Copy runs with formatting
                                for run in para.runs:
                                    new_run = cell_para.add_run(run.text)
                                    new_run.bold = run.bold
                                    new_run.italic = run.italic
                                    new_run.underline = run.underline
                                    if run.font.size:
                                        new_run.font.size = run.font.size
                                    if run.font.name:
                                        new_run.font.name = run.font.name
                    return True  # Success

            return False  # Table not found

        except Exception as e:
            print(f"Fallback table copy failed: {e}")
            return False

    def create_translator_certificate(self, output_file, metadata):
        """Create translator certificate"""

        doc = Document()

        # Configure section - no header/footer, Portrait orientation, Letter size
        section = doc.sections[0]
        section.page_height = Inches(11)  # Letter height
        section.page_width = Inches(8.5)  # Letter width
        section.orientation = 0  # 0 = Portrait, 1 = Landscape
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # Clear header and footer completely
        for paragraph in section.header.paragraphs:
            paragraph.clear()
        for paragraph in section.footer.paragraphs:
            paragraph.clear()

        # Add logo (left-aligned)
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Try to add logo image
        logo_path = self.assets_dir / 'park_logo.png'
        logo_added = False

        if logo_path.exists():
            try:
                run = logo_para.add_run()
                run.add_picture(str(logo_path), height=Inches(0.8))
                logo_added = True
                print(f"Successfully added Park logo to translator certificate from: {logo_path}")
            except Exception as e:
                print(f"Error adding logo image to translator certificate: {e}")
                import traceback
                traceback.print_exc()

        # If logo couldn't be loaded, use text fallback
        if not logo_added:
            run = logo_para.add_run("PARK EVALUATION SERVICES")
            run.font.size = Pt(14)
            run.font.bold = True
            print(f"Warning: Park logo file not found at: {logo_path}")

        # Add DATE at top right (with actual date)
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = date_para.add_run(metadata['date'])
        run.font.size = Pt(11)

        # Title (no extra spacing before)
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("TRANSLATION CERTIFICATION")
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.all_caps = True

        # Certificate text - First paragraph (justified) (no extra spacing before)
        para1_text = f"I, {metadata['translator_name']}, am fluent and competent in both the {metadata['source_language']} and {metadata['target_language']} languages. I certify that the enclosed translation is true, complete, and accurate and that, to the best of my knowledge and belief, the translation accurately reflects the meaning and intention of the original text. I am not a family member, friend, or business associate of anyone referenced in this translation but a completely disinterested third party with no relationship to the beneficiary."

        para1 = doc.add_paragraph(para1_text)
        para1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in para1.runs:
            run.font.size = Pt(11)

        doc.add_paragraph()  # Spacing

        # Certificate text - Second paragraph (justified)
        para2_text = "This is to certify the correctness of the translation only. I do not make any claims or guarantees about the authenticity or content of the original document."

        para2 = doc.add_paragraph(para2_text)
        para2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in para2.runs:
            run.font.size = Pt(11)

        doc.add_paragraph()  # Spacing

        # Add "Sincerely,"
        sincerely = doc.add_paragraph("Sincerely,")
        sincerely.alignment = WD_ALIGN_PARAGRAPH.LEFT
        sincerely.runs[0].font.size = Pt(11)

        # Add signature image if available
        sig_added = False
        if metadata.get('translator_signature'):
            sig_path = metadata['translator_signature']
            print(f"Looking for translator signature at: {sig_path}")

            if os.path.exists(sig_path):
                try:
                    sig_img_para = doc.add_paragraph()
                    run = sig_img_para.add_run()
                    run.add_picture(sig_path, width=Inches(2))
                    sig_img_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    sig_added = True
                    print(f"Successfully added translator signature from: {sig_path}")
                except Exception as e:
                    print(f"Error adding signature image: {e}")
                    import traceback
                    traceback.print_exc()
            else:
                print(f"Warning: Translator signature file not found: {sig_path}")

        if not sig_added:
            # Add placeholder for signature
            doc.add_paragraph()  # Spacing

        # Signature line
        sig_para = doc.add_paragraph("_" * 40)
        sig_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Translator name under signature line
        sig_label = doc.add_paragraph(metadata['translator_name'])
        sig_label.alignment = WD_ALIGN_PARAGRAPH.LEFT
        sig_label.runs[0].font.size = Pt(11)

        # Add actual case number
        case_para = doc.add_paragraph(f"Case Number: {metadata['case_number']}")
        case_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        case_para.runs[0].font.size = Pt(11)

        doc.add_paragraph()  # Spacing
        doc.add_paragraph()  # Additional spacing
        doc.add_paragraph()  # Additional spacing
        doc.add_paragraph()  # Additional spacing

        # Footer with contact info image
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Try to add contact info image
        contact_img_path = self.assets_dir / 'park_contact.png'
        contact_img_added = False

        if contact_img_path.exists():
            try:
                run = footer_para.add_run()
                run.add_picture(str(contact_img_path), width=Inches(4.375))  # 3.5 * 1.25 = 25% larger
                contact_img_added = True
                print(f"Successfully added contact info image from: {contact_img_path}")
            except Exception as e:
                print(f"Error adding contact info image: {e}")
                import traceback
                traceback.print_exc()

        # Fallback to text if image not available
        if not contact_img_added:
            run = footer_para.add_run("(212) 581-8877 • www.ParkEval.com")
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(97, 145, 43)
            print(f"Warning: Contact info image not found at: {contact_img_path}")

        # Save certificate
        doc.save(output_file)
        print(f"Saved translator certificate to: {output_file}")

    def create_park_certificate(self, output_file, metadata):
        """Create Park employee certificate"""

        doc = Document()

        # Configure section - no header/footer, Portrait orientation, Letter size
        section = doc.sections[0]
        section.page_height = Inches(11)  # Letter height
        section.page_width = Inches(8.5)  # Letter width
        section.orientation = 0  # 0 = Portrait, 1 = Landscape
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # Clear header and footer completely
        for paragraph in section.header.paragraphs:
            paragraph.clear()
        for paragraph in section.footer.paragraphs:
            paragraph.clear()

        # Add logo (left-aligned)
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Try to add logo image (use EEI logo for EEI division, Park logo otherwise)
        division = metadata.get('division', 'PARK')
        if division == 'EEI':
            logo_path = self.assets_dir / 'eei_logo.png'
            logo_name = "EEI"
            fallback_text = "EEI"
        else:
            logo_path = self.assets_dir / 'park_logo.png'
            logo_name = "Park"
            fallback_text = "PARK EVALUATION SERVICES"

        logo_added = False

        if logo_path.exists():
            try:
                run = logo_para.add_run()
                run.add_picture(str(logo_path), height=Inches(0.8))
                logo_added = True
                print(f"Successfully added {logo_name} logo to {logo_name} certificate from: {logo_path}")
            except Exception as e:
                print(f"Error adding logo image to {logo_name} certificate: {e}")
                import traceback
                traceback.print_exc()

        # If logo couldn't be loaded, use text fallback
        if not logo_added:
            run = logo_para.add_run(fallback_text)
            run.font.size = Pt(14)
            run.font.bold = True
            print(f"Warning: {logo_name} logo file not found at: {logo_path}")

        # Add DATE at top right (with actual date)
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = date_para.add_run(metadata['date'])
        run.font.size = Pt(11)

        # Title
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("TRANSLATION CERTIFICATION")
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.all_caps = True

        doc.add_paragraph()  # Spacing

        # Certificate text - First paragraph (justified)
        para1_text = f"This is to certify that the enclosed {metadata['source_language']} to {metadata['target_language']} translation (Case Number: {metadata['case_number']}) was made under my personal supervision by a qualified translator who is fluent in both languages, and that to the best of my knowledge and understanding is a true and complete rendition of the corresponding original document. This document has not been translated for a family member, friend, or business associate but by a completely disinterested third party with no relationship to the beneficiary."

        para1 = doc.add_paragraph(para1_text)
        para1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in para1.runs:
            run.font.size = Pt(11)

        doc.add_paragraph()  # Spacing

        # Certificate text - Second paragraph (justified)
        para2_text = "This is to certify the correctness of the translation only. We do not make any claims or guarantees about the authenticity or content of the original document. Further, Park Evaluations assumes no liability for the way in which the translation is used by the customer or any third party, including end-users of the translation."

        para2 = doc.add_paragraph(para2_text)
        para2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in para2.runs:
            run.font.size = Pt(11)

        doc.add_paragraph()  # Spacing

        # Signature section (Park employee)
        sig_para = doc.add_paragraph("_" * 40)
        sig_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Name line
        sig_name = doc.add_paragraph("Howard Borenstein")
        sig_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
        sig_name.runs[0].font.size = Pt(11)

        # Title line
        sig_title = doc.add_paragraph("Evaluator")
        sig_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        sig_title.runs[0].font.size = Pt(11)

        # Add Park signature image from assets directory
        park_sig_path = self.assets_dir / 'park_signature.png'
        sig_added = False

        print(f"Looking for Park signature at: {park_sig_path}")

        if park_sig_path.exists():
            try:
                # Insert signature image ABOVE the line
                run = sig_para.insert_paragraph_before().add_run()
                run.add_picture(str(park_sig_path), width=Inches(2))
                sig_added = True
                print(f"Successfully added Park signature from: {park_sig_path}")
            except Exception as e:
                print(f"Error adding Park signature image: {e}")
                import traceback
                traceback.print_exc()
        else:
            print(f"Warning: Park signature file not found: {park_sig_path}")

        if not sig_added:
            # Add placeholder for signature
            doc.add_paragraph()  # Spacing

        doc.add_paragraph()  # Spacing
        doc.add_paragraph()  # Spacing
        doc.add_paragraph()  # Additional spacing
        doc.add_paragraph()  # Additional spacing
        doc.add_paragraph()  # Additional spacing

        # Footer with contact info image - skip for EEI division
        division = metadata.get('division', 'PARK')
        if division != 'EEI':
            footer_para = doc.add_paragraph()
            footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Try to add contact info image
            contact_img_path = self.assets_dir / 'park_contact.png'
            contact_img_added = False

            if contact_img_path.exists():
                try:
                    run = footer_para.add_run()
                    run.add_picture(str(contact_img_path), width=Inches(4.375))  # 3.5 * 1.25 = 25% larger
                    contact_img_added = True
                    print(f"Successfully added contact info image from: {contact_img_path}")
                except Exception as e:
                    print(f"Error adding contact info image: {e}")
                    import traceback
                    traceback.print_exc()

            # Fallback to text if image not available
            if not contact_img_added:
                run = footer_para.add_run("(212) 581-8877 • www.ParkEval.com")
                run.font.name = 'Arial'
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(97, 145, 43)
                print(f"Warning: Contact info image not found at: {contact_img_path}")
        else:
            print("Skipping contact info image (EEI division)")

        # Save certificate
        doc.save(output_file)
        print(f"Saved Park certificate to: {output_file}")

    def combine_documents(self, doc_paths, output_path):
        """Combine multiple Word documents into one using safe section-based approach"""
        from io import BytesIO
        import zipfile
        from lxml import etree

        # Start with the first document (translation with header/footer)
        combined = Document(doc_paths[0])

        # Set the first section to restart page numbering at 1
        # This ensures page numbers in the translation section don't include certificate pages
        first_section = combined.sections[0]
        sectPr = first_section._sectPr
        pgNumType = sectPr.find(qn('w:pgNumType'))
        if pgNumType is None:
            pgNumType = OxmlElement('w:pgNumType')
            sectPr.append(pgNumType)
        pgNumType.set(qn('w:start'), '1')

        # Add remaining documents (certificates)
        for idx, doc_path in enumerate(doc_paths[1:], start=1):
            # Read the document to append
            sub_doc = Document(doc_path)

            # Add a new section for this document (creates page break and allows different header/footer)
            new_section = combined.add_section()

            # Set standard Letter size and Portrait orientation for certificates
            # This prevents certificates from inheriting custom sizing from source documents
            new_section.page_height = Inches(11)  # Letter height
            new_section.page_width = Inches(8.5)  # Letter width
            new_section.orientation = 0  # 0 = Portrait, 1 = Landscape

            # Force single column layout for certificates
            # Explicitly set to 1 column (don't just remove - must define it)
            sectPr = new_section._sectPr
            cols_elem = sectPr.find(qn('w:cols'))
            if cols_elem is None:
                cols_elem = OxmlElement('w:cols')
                sectPr.append(cols_elem)
            # Set to single column explicitly
            cols_elem.set(qn('w:num'), '1')

            # Make sure this section doesn't inherit header/footer from previous section
            new_section.header.is_linked_to_previous = False
            new_section.footer.is_linked_to_previous = False

            # Clear header and footer for this section (certificates shouldn't have them)
            for paragraph in new_section.header.paragraphs:
                paragraph.clear()
            for paragraph in new_section.footer.paragraphs:
                paragraph.clear()

            # Restart page numbering for this section at 1
            # This ensures certificates don't continue the page count from translation pages
            sectPr = new_section._sectPr
            pgNumType = sectPr.find(qn('w:pgNumType'))
            if pgNumType is None:
                pgNumType = OxmlElement('w:pgNumType')
                sectPr.append(pgNumType)
            pgNumType.set(qn('w:start'), '1')

            # Copy all paragraphs from sub_doc
            for paragraph in sub_doc.paragraphs:
                # Create a new paragraph in the combined document
                new_para = combined.add_paragraph()
                new_para.alignment = paragraph.alignment
                new_para.style = paragraph.style

                # Copy paragraph formatting
                if paragraph.runs:
                    for run in paragraph.runs:
                        # Check if run contains an image (drawing or picture)
                        has_image = False

                        # Check for drawing elements (modern Word images)
                        drawings = run._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing', run._element.nsmap if hasattr(run._element, 'nsmap') else None)

                        if not drawings:
                            # Also check without namespace for older formats
                            for child in run._element:
                                if 'drawing' in child.tag or 'pict' in child.tag:
                                    drawings = [child]
                                    break

                        if drawings:
                            has_image = True
                            # Try to extract and copy the image
                            for drawing in drawings:
                                try:
                                    # Find the blip element that contains the image relationship ID
                                    # Try multiple ways to find it
                                    blip = None

                                    # Method 1: Use find with full namespace
                                    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                                          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
                                    blip = drawing.find('.//a:blip', ns)

                                    # Method 2: Search all descendants
                                    if blip is None:
                                        for elem in drawing.iter():
                                            if 'blip' in elem.tag.lower():
                                                blip = elem
                                                break

                                    if blip is not None:
                                        # Get the relationship ID
                                        rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                        if not rId:
                                            # Try without namespace
                                            rId = blip.get('embed')

                                        if rId and rId in sub_doc.part.related_parts:
                                            # Get the image part
                                            image_part = sub_doc.part.related_parts[rId]
                                            image_bytes = image_part.blob

                                            # Try to get image dimensions from the extent element
                                            extent = drawing.find('.//wp:extent', {'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'})
                                            if extent is None:
                                                # Try searching all descendants
                                                for elem in drawing.iter():
                                                    if 'extent' in elem.tag.lower():
                                                        extent = elem
                                                        break

                                            # Add the image to the new paragraph
                                            new_run = new_para.add_run()

                                            if extent is not None and extent.get('cx'):
                                                # Convert EMU to inches (1 inch = 914400 EMUs)
                                                try:
                                                    cx = int(extent.get('cx'))
                                                    width_inches = cx / 914400.0
                                                    new_run.add_picture(BytesIO(image_bytes), width=Inches(width_inches))
                                                except:
                                                    # Fallback to default size
                                                    new_run.add_picture(BytesIO(image_bytes))
                                            else:
                                                # No size info, use default
                                                new_run.add_picture(BytesIO(image_bytes))

                                            print(f"Successfully copied image from certificate")
                                        else:
                                            print(f"Warning: Could not find image relationship {rId}")
                                            new_run = new_para.add_run("[Image not found]")
                                    else:
                                        print(f"Warning: Could not find blip element in drawing")
                                        new_run = new_para.add_run("[Image]")

                                except Exception as e:
                                    print(f"Warning: Could not copy image: {e}")
                                    import traceback
                                    traceback.print_exc()
                                    new_run = new_para.add_run("[Image error]")

                        # If no image, copy text and formatting
                        if not has_image:
                            new_run = new_para.add_run(run.text)
                            new_run.bold = run.bold
                            new_run.italic = run.italic
                            new_run.underline = run.underline
                            if run.font.size:
                                new_run.font.size = run.font.size
                            if run.font.name:
                                new_run.font.name = run.font.name

            # Copy tables
            for table in sub_doc.tables:
                new_table = combined.add_table(rows=len(table.rows), cols=len(table.columns))
                for i, row in enumerate(table.rows):
                    for j, cell in enumerate(row.cells):
                        new_table.rows[i].cells[j].text = cell.text

        # Save combined document
        combined.save(output_path)
        print(f"Saved combined document to: {output_path}")

    def convert_to_pdf(self, docx_path, pdf_path):
        """Convert Word document to PDF"""

        try:
            # Try using docx2pdf (works on Windows with Word installed)
            # Initialize COM for the current thread
            try:
                import pythoncom
                pythoncom.CoInitialize()
                com_initialized = True
            except:
                com_initialized = False

            try:
                from docx2pdf import convert
                convert(str(docx_path), str(pdf_path))
                print(f"Converted to PDF using docx2pdf: {pdf_path}")
                return True
            finally:
                # Uninitialize COM
                if com_initialized:
                    try:
                        pythoncom.CoUninitialize()
                    except:
                        pass

        except Exception as e:
            print(f"docx2pdf conversion failed: {e}")

            # Fallback: Try using win32com (Windows only)
            try:
                import pythoncom
                import win32com.client

                # Initialize COM for the current thread
                pythoncom.CoInitialize()

                try:
                    word = win32com.client.Dispatch("Word.Application")
                    word.Visible = False

                    doc = word.Documents.Open(str(Path(docx_path).absolute()))
                    doc.SaveAs(str(Path(pdf_path).absolute()), FileFormat=17)  # 17 = PDF
                    doc.Close()
                    word.Quit()

                    print(f"Converted to PDF using Word COM: {pdf_path}")
                    return True
                finally:
                    # Uninitialize COM
                    try:
                        pythoncom.CoUninitialize()
                    except:
                        pass

            except Exception as e2:
                print(f"Word COM conversion also failed: {e2}")

                # If both methods fail, save as DOCX with note
                fallback_path = str(pdf_path).replace('.pdf', '_CONVERT_TO_PDF.docx')
                shutil.copy(docx_path, fallback_path)
                print(f"Could not convert to PDF. Saved as Word document: {fallback_path}")
                print("Please manually convert to PDF using Microsoft Word.")
                return False
